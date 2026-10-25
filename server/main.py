"""성취평가 도우미 서버 — 데이터는 data/assessment.db (SQLite) 하나에 저장한다.

처음 실행해 표가 비어 있으면 data/seed/*.json 을 넣는다.
"""
import json
import sqlite3
from contextlib import closing
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / "data" / "assessment.db"
SEED_DIR = ROOT / "data" / "seed"
ANALYSIS_SEED_DIR = ROOT / "data" / "seed-analysis"  # 성취기준 분석 초안 (school_level·subject·code 로 성취기준과 이음)
CURRICULUM_PATH = ROOT / "data" / "curriculum.json"  # 교육과정 원문(과목·영역 정보) — 읽기 전용
REVIEW_TEMPLATE_PATH = ROOT / "data" / "review-template.json"  # 평가계획서 검토 항목(검토 템플릿 xlsx 에서 뽑음)
REVIEW_XLSX_PATH = ROOT / "data" / "templates" / "평가계획서 검토 템플릿.xlsx"  # 결과 내려받기 양식
PLANS_DIR = ROOT / "data" / "plans"  # 올라온 평가계획서 원본
PLAN_EXTS = (".pdf", ".hwp", ".hwpx")
PLAN_MAX_BYTES = 30 * 1024 * 1024

LEVELS = ("A", "B", "C", "D", "E")
# 지식·이해 / 과정·기능 / 가치·태도 / 평가요소 — 줄마다 한 항목인 목록, JSON 문자열로 저장
LIST_FIELDS = ("knowledge", "skills", "attitude", "elements")

SCHEMA = """
CREATE TABLE IF NOT EXISTS standards (
  id           INTEGER PRIMARY KEY,
  school_level TEXT NOT NULL,
  subject      TEXT NOT NULL,
  code         TEXT NOT NULL,
  statement    TEXT NOT NULL,
  explanation  TEXT NOT NULL DEFAULT '',
  level_a TEXT NOT NULL DEFAULT '', level_b TEXT NOT NULL DEFAULT '', level_c TEXT NOT NULL DEFAULT '',
  level_d TEXT NOT NULL DEFAULT '', level_e TEXT NOT NULL DEFAULT '',
  knowledge TEXT NOT NULL DEFAULT '[]',
  skills    TEXT NOT NULL DEFAULT '[]',
  attitude  TEXT NOT NULL DEFAULT '[]',
  elements  TEXT NOT NULL DEFAULT '[]',
  ai_draft  INTEGER NOT NULL DEFAULT 0,
  source    TEXT NOT NULL DEFAULT '',
  updated_at TEXT NOT NULL,
  UNIQUE (school_level, subject, code)
);
-- 성취기준 분석(연수 양식 ②~⑦). 항목이 많고 자주 바뀌어 JSON 하나로 둔다.
CREATE TABLE IF NOT EXISTS analyses (
  standard_id INTEGER PRIMARY KEY REFERENCES standards(id),
  data        TEXT NOT NULL,
  ai_draft    INTEGER NOT NULL DEFAULT 0,
  updated_at  TEXT NOT NULL
);
-- 문항 개발(수행평가 과제 및 채점기준표 양식). 양식 칸이 많아 JSON 하나로 둔다.
-- status: draft(작성 중) | requested(Claude Code 초안 요청) · ai_draft: Claude Code 가 채운 초안인지
CREATE TABLE IF NOT EXISTS item_sets (
  id           INTEGER PRIMARY KEY,
  school_level TEXT NOT NULL,
  subject      TEXT NOT NULL,
  data         TEXT NOT NULL,
  status       TEXT NOT NULL DEFAULT 'draft',
  ai_draft     INTEGER NOT NULL DEFAULT 0,
  updated_at   TEXT NOT NULL
);
-- 평가계획서 검토. 올리면 pending, Claude Code 가 검토 결과(result JSON)를 넣으면 done.
CREATE TABLE IF NOT EXISTS plan_reviews (
  id          INTEGER PRIMARY KEY,
  filename    TEXT NOT NULL,
  stored      TEXT NOT NULL,
  status      TEXT NOT NULL DEFAULT 'pending',
  result      TEXT,
  uploaded_at TEXT NOT NULL,
  reviewed_at TEXT
);
"""


class Standard(BaseModel):
    school_level: str
    subject: str
    code: str
    statement: str
    explanation: str = ""  # 성취기준 해설
    levels: dict[str, str] = {}
    knowledge: list[str] = []
    skills: list[str] = []
    attitude: list[str] = []
    elements: list[str] = []
    ai_draft: bool = False
    source: str = ""


def connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def to_row(s: Standard) -> dict:
    row = {
        "school_level": s.school_level.strip(),
        "subject": s.subject.strip(),
        "code": s.code.strip(),
        "statement": s.statement.strip(),
        "explanation": s.explanation.strip(),
        "ai_draft": int(s.ai_draft),
        "source": s.source,
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    for lv in LEVELS:
        row[f"level_{lv.lower()}"] = s.levels.get(lv, "").strip()
    for f in LIST_FIELDS:
        row[f] = json.dumps([x.strip() for x in getattr(s, f) if x.strip()], ensure_ascii=False)
    return row


def from_row(r: sqlite3.Row) -> dict:
    d = {k: r[k] for k in ("id", "school_level", "subject", "code", "statement", "explanation", "source", "updated_at")}
    d["levels"] = {lv: r[f"level_{lv.lower()}"] for lv in LEVELS}
    for f in LIST_FIELDS:
        d[f] = json.loads(r[f])
    d["ai_draft"] = bool(r["ai_draft"])
    return d


def insert(conn: sqlite3.Connection, s: Standard) -> int:
    row = to_row(s)
    cols = ", ".join(row)
    marks = ", ".join(f":{k}" for k in row)
    return conn.execute(f"INSERT INTO standards ({cols}) VALUES ({marks})", row).lastrowid


def seed_analyses(conn: sqlite3.Connection) -> int:
    """data/seed-analysis/*.json 중 성취기준은 있고 분석은 아직 없는 것만 넣는다. 넣은 개수를 돌려준다."""
    added = 0
    for path in sorted(ANALYSIS_SEED_DIR.glob("*.json")):
        for item in json.loads(path.read_text(encoding="utf-8")):
            row = conn.execute(
                "SELECT id FROM standards WHERE school_level = ? AND subject = ? AND code = ?",
                (item["school_level"], item["subject"], item["code"]),
            ).fetchone()
            if not row or conn.execute("SELECT 1 FROM analyses WHERE standard_id = ?", (row[0],)).fetchone():
                continue
            conn.execute(
                "INSERT INTO analyses (standard_id, data, ai_draft, updated_at) VALUES (?, ?, ?, ?)",
                (row[0], json.dumps(item["data"], ensure_ascii=False), int(item.get("ai_draft", True)),
                 datetime.now().isoformat(timespec="seconds")),
            )
            added += 1
    return added


def init_db() -> None:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with closing(connect()) as conn, conn:
        conn.executescript(SCHEMA)
        # 해설 열이 생기기 전에 만든 DB
        if "explanation" not in [c[1] for c in conn.execute("PRAGMA table_info(standards)")]:
            conn.execute("ALTER TABLE standards ADD COLUMN explanation TEXT NOT NULL DEFAULT ''")
        if conn.execute("SELECT COUNT(*) FROM standards").fetchone()[0]:
            return
        for path in sorted(SEED_DIR.glob("*.json")):
            for item in json.loads(path.read_text(encoding="utf-8")):
                insert(conn, Standard(**item))
        seed_analyses(conn)


init_db()
app = FastAPI(title="성취평가 도우미")


@app.get("/api/curriculum")
def get_curriculum(school_level: str, subject: str) -> dict | None:
    """과목의 성격·목표·교수학습·평가와 영역별 내용 체계·고려 사항. 없으면 null."""
    if not CURRICULUM_PATH.exists():
        return None
    for c in json.loads(CURRICULUM_PATH.read_text(encoding="utf-8")):
        if c["school_level"] == school_level and c["subject"] == subject:
            return c
    return None


@app.get("/api/standards")
def list_standards(school_level: str | None = None, subject: str | None = None) -> list[dict]:
    sql, args = "SELECT * FROM standards WHERE 1=1", []
    if school_level:
        sql += " AND school_level = ?"
        args.append(school_level)
    if subject:
        sql += " AND subject = ?"
        args.append(subject)
    with closing(connect()) as conn:
        return [from_row(r) for r in conn.execute(sql + " ORDER BY code", args)]


@app.post("/api/standards")
def create_standard(s: Standard) -> dict:
    with closing(connect()) as conn, conn:
        try:
            new_id = insert(conn, s)
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"{s.code} 성취기준이 이미 있습니다.")
        return from_row(conn.execute("SELECT * FROM standards WHERE id = ?", (new_id,)).fetchone())


@app.put("/api/standards/{sid}")
def update_standard(sid: int, s: Standard) -> dict:
    row = to_row(s)
    sets = ", ".join(f"{k} = :{k}" for k in row)
    with closing(connect()) as conn, conn:
        try:
            cur = conn.execute(f"UPDATE standards SET {sets} WHERE id = :id", {**row, "id": sid})
        except sqlite3.IntegrityError:
            raise HTTPException(409, f"{s.code} 성취기준이 이미 있습니다.")
        if not cur.rowcount:
            raise HTTPException(404, "성취기준을 찾을 수 없습니다.")
        return from_row(conn.execute("SELECT * FROM standards WHERE id = ?", (sid,)).fetchone())


@app.delete("/api/standards/{sid}")
def delete_standard(sid: int) -> dict:
    with closing(connect()) as conn, conn:
        conn.execute("DELETE FROM analyses WHERE standard_id = ?", (sid,))
        if not conn.execute("DELETE FROM standards WHERE id = ?", (sid,)).rowcount:
            raise HTTPException(404, "성취기준을 찾을 수 없습니다.")
    return {"ok": True}


class Analysis(BaseModel):
    data: dict
    ai_draft: bool = False


@app.get("/api/analyses/{sid}")
def get_analysis(sid: int) -> dict | None:
    """성취기준 분석. 아직 없으면 null."""
    with closing(connect()) as conn:
        r = conn.execute("SELECT * FROM analyses WHERE standard_id = ?", (sid,)).fetchone()
    if not r:
        return None
    return {"standard_id": sid, "data": json.loads(r["data"]), "ai_draft": bool(r["ai_draft"]), "updated_at": r["updated_at"]}


@app.put("/api/analyses/{sid}")
def save_analysis(sid: int, a: Analysis) -> dict:
    now = datetime.now().isoformat(timespec="seconds")
    with closing(connect()) as conn, conn:
        if not conn.execute("SELECT 1 FROM standards WHERE id = ?", (sid,)).fetchone():
            raise HTTPException(404, "성취기준을 찾을 수 없습니다.")
        conn.execute(
            "INSERT INTO analyses (standard_id, data, ai_draft, updated_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(standard_id) DO UPDATE SET data = excluded.data, ai_draft = excluded.ai_draft, updated_at = excluded.updated_at",
            (sid, json.dumps(a.data, ensure_ascii=False), int(a.ai_draft), now),
        )
    return get_analysis(sid)


# ── 평가계획서 검토 ─────────────────────────────────────────────
# 올리기·목록·결과 보기·엑셀 내려받기는 앱이, 검토(result 채우기)는 Claude Code 가 한다(docs/평가계획서-검토-절차.md).

VERDICTS = {"ok": "✔ 잘 됨", "warn": "△ 고민 필요", "bad": "✖ 문제", "na": "－ 해당 없음"}


class ReviewItem(BaseModel):
    verdict: str  # ok | warn | bad | na
    where: str = ""  # 계획서에서 찾은 위치
    finding: str = ""  # 판단 근거 (무엇이 어떻게 되어 있나)
    fix: str = ""  # 수정 제안


class ReviewResult(BaseModel):
    summary: str
    items: dict[str, ReviewItem]  # 검토 항목 ID(가1, 1 …) → 판정


def review_row(r: sqlite3.Row) -> dict:
    return {
        "id": r["id"],
        "filename": r["filename"],
        "status": r["status"],
        "result": json.loads(r["result"]) if r["result"] else None,
        "uploaded_at": r["uploaded_at"],
        "reviewed_at": r["reviewed_at"],
    }


def get_review_row(conn: sqlite3.Connection, rid: int) -> sqlite3.Row:
    r = conn.execute("SELECT * FROM plan_reviews WHERE id = ?", (rid,)).fetchone()
    if not r:
        raise HTTPException(404, "평가계획서를 찾을 수 없습니다.")
    return r


@app.get("/api/review-template")
def get_review_template() -> list[dict]:
    return json.loads(REVIEW_TEMPLATE_PATH.read_text(encoding="utf-8"))


@app.get("/api/plan-reviews")
def list_plan_reviews() -> list[dict]:
    with closing(connect()) as conn:
        return [review_row(r) for r in conn.execute("SELECT * FROM plan_reviews ORDER BY id DESC")]


@app.post("/api/plan-reviews")
async def upload_plan(filename: str, request: Request) -> dict:
    """본문에 파일 바이트를 그대로 받는다(multipart 라이브러리 없이)."""
    name = Path(filename).name.strip()
    ext = Path(name).suffix.lower()
    if ext not in PLAN_EXTS:
        raise HTTPException(400, "PDF, HWP, HWPX 파일만 올릴 수 있어요.")
    body = await request.body()
    if not body:
        raise HTTPException(400, "빈 파일이에요.")
    if len(body) > PLAN_MAX_BYTES:
        raise HTTPException(400, "파일이 너무 커요(30MB 이하).")
    PLANS_DIR.mkdir(parents=True, exist_ok=True)
    now = datetime.now().isoformat(timespec="seconds")
    with closing(connect()) as conn, conn:
        rid = conn.execute(
            "INSERT INTO plan_reviews (filename, stored, uploaded_at) VALUES (?, '', ?)", (name, now)
        ).lastrowid
        stored = f"{rid}{ext}"
        (PLANS_DIR / stored).write_bytes(body)
        conn.execute("UPDATE plan_reviews SET stored = ? WHERE id = ?", (stored, rid))
        return review_row(get_review_row(conn, rid))


@app.get("/api/plan-reviews/{rid}")
def get_plan_review(rid: int) -> dict:
    with closing(connect()) as conn:
        return review_row(get_review_row(conn, rid))


@app.get("/api/plan-reviews/{rid}/file")
def get_plan_file(rid: int) -> FileResponse:
    with closing(connect()) as conn:
        r = get_review_row(conn, rid)
    return FileResponse(PLANS_DIR / r["stored"], filename=r["filename"])


@app.put("/api/plan-reviews/{rid}/result")
def save_plan_result(rid: int, res: ReviewResult) -> dict:
    """Claude Code 가 검토를 마치고 결과를 넣는다."""
    bad = [k for k, v in res.items.items() if v.verdict not in VERDICTS]
    if bad:
        raise HTTPException(400, f"판정 값이 잘못된 항목: {', '.join(bad)}")
    now = datetime.now().isoformat(timespec="seconds")
    with closing(connect()) as conn, conn:
        get_review_row(conn, rid)
        conn.execute(
            "UPDATE plan_reviews SET status = 'done', result = ?, reviewed_at = ? WHERE id = ?",
            (res.model_dump_json(), now, rid),
        )
        return review_row(get_review_row(conn, rid))


@app.delete("/api/plan-reviews/{rid}")
def delete_plan_review(rid: int) -> dict:
    with closing(connect()) as conn, conn:
        r = get_review_row(conn, rid)
        conn.execute("DELETE FROM plan_reviews WHERE id = ?", (rid,))
    (PLANS_DIR / r["stored"]).unlink(missing_ok=True)
    return {"ok": True}


@app.get("/api/plan-reviews/{rid}/xlsx")
def download_plan_review(rid: int) -> Response:
    """검토 템플릿 양식에 판정·근거·수정 제안 열을 더해 내려준다."""
    from copy import copy
    from io import BytesIO
    from urllib.parse import quote

    import openpyxl

    with closing(connect()) as conn:
        r = review_row(get_review_row(conn, rid))
    if not r["result"]:
        raise HTTPException(409, "아직 검토가 끝나지 않았어요.")
    items = r["result"]["items"]

    wb = openpyxl.load_workbook(REVIEW_XLSX_PATH)
    # 시트마다: 결과를 쓸 첫 열(판정), 그 뒤로 근거·수정 제안
    for ws, first_col in zip(wb.worksheets, (6, 7)):
        head = ws.cell(row=4, column=5)
        for i, title in enumerate(("판정", "찾은 위치 · 근거", "수정 제안")):
            c = ws.cell(row=4, column=first_col + i, value=title)
            c.font, c.fill, c.border, c.alignment = copy(head.font), copy(head.fill), copy(head.border), copy(head.alignment)
        ws.column_dimensions[openpyxl.utils.get_column_letter(first_col)].width = 14
        ws.column_dimensions[openpyxl.utils.get_column_letter(first_col + 1)].width = 60
        ws.column_dimensions[openpyxl.utils.get_column_letter(first_col + 2)].width = 50
        for row in range(5, ws.max_row + 1):
            key = ws.cell(row=row, column=1).value
            it = items.get(str(key)) if key is not None else None
            if not it:
                continue
            body = ws.cell(row=row, column=4)
            where = f"[{it['where']}] " if it.get("where") else ""
            for i, val in enumerate((VERDICTS[it["verdict"]], where + it.get("finding", ""), it.get("fix", ""))):
                c = ws.cell(row=row, column=first_col + i, value=val)
                c.font, c.border, c.alignment = copy(body.font), copy(body.border), copy(body.alignment)

    s = wb.create_sheet("검토 요약", 0)
    s["A1"], s["A2"], s["A3"] = "평가계획서 검토 결과", f"파일: {r['filename']}", f"검토일: {r['reviewed_at']}"
    s["A5"] = r["result"]["summary"]
    s["A1"].font = openpyxl.styles.Font(bold=True, size=14)
    s.column_dimensions["A"].width = 110
    s["A5"].alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    buf = BytesIO()
    wb.save(buf)
    name = f"{Path(r['filename']).stem} 검토 결과.xlsx"
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )


# ── 문항 개발 ─────────────────────────────────────────────
# 입력·저장·워드 내려받기는 앱이, 초안 채우기는 Claude Code 가 한다(docs/문항-초안-절차.md).


class ItemSet(BaseModel):
    school_level: str
    subject: str
    data: dict
    status: str = "draft"  # draft | requested
    ai_draft: bool = False


def item_row(r: sqlite3.Row) -> dict:
    return {
        "id": r["id"],
        "school_level": r["school_level"],
        "subject": r["subject"],
        "data": json.loads(r["data"]),
        "status": r["status"],
        "ai_draft": bool(r["ai_draft"]),
        "updated_at": r["updated_at"],
    }


def get_item_row(conn: sqlite3.Connection, iid: int) -> sqlite3.Row:
    r = conn.execute("SELECT * FROM item_sets WHERE id = ?", (iid,)).fetchone()
    if not r:
        raise HTTPException(404, "문항을 찾을 수 없습니다.")
    return r


def check_status(status: str) -> None:
    if status not in ("draft", "requested"):
        raise HTTPException(400, "status 는 draft 또는 requested 여야 합니다.")


@app.get("/api/item-sets")
def list_item_sets(school_level: str | None = None, subject: str | None = None) -> list[dict]:
    sql, args = "SELECT * FROM item_sets WHERE 1=1", []
    if school_level:
        sql += " AND school_level = ?"
        args.append(school_level)
    if subject:
        sql += " AND subject = ?"
        args.append(subject)
    with closing(connect()) as conn:
        return [item_row(r) for r in conn.execute(sql + " ORDER BY id DESC", args)]


@app.post("/api/item-sets")
def create_item_set(it: ItemSet) -> dict:
    check_status(it.status)
    now = datetime.now().isoformat(timespec="seconds")
    with closing(connect()) as conn, conn:
        iid = conn.execute(
            "INSERT INTO item_sets (school_level, subject, data, status, ai_draft, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
            (it.school_level, it.subject, json.dumps(it.data, ensure_ascii=False), it.status, int(it.ai_draft), now),
        ).lastrowid
        return item_row(get_item_row(conn, iid))


@app.get("/api/item-sets/{iid}")
def get_item_set(iid: int) -> dict:
    with closing(connect()) as conn:
        return item_row(get_item_row(conn, iid))


@app.put("/api/item-sets/{iid}")
def save_item_set(iid: int, it: ItemSet) -> dict:
    check_status(it.status)
    now = datetime.now().isoformat(timespec="seconds")
    with closing(connect()) as conn, conn:
        get_item_row(conn, iid)
        conn.execute(
            "UPDATE item_sets SET school_level = ?, subject = ?, data = ?, status = ?, ai_draft = ?, updated_at = ? WHERE id = ?",
            (it.school_level, it.subject, json.dumps(it.data, ensure_ascii=False), it.status, int(it.ai_draft), now, iid),
        )
        return item_row(get_item_row(conn, iid))


@app.delete("/api/item-sets/{iid}")
def delete_item_set(iid: int) -> dict:
    with closing(connect()) as conn, conn:
        if not conn.execute("DELETE FROM item_sets WHERE id = ?", (iid,)).rowcount:
            raise HTTPException(404, "문항을 찾을 수 없습니다.")
    return {"ok": True}


@app.get("/api/item-sets/{iid}/docx")
def download_item_set(iid: int) -> Response:
    """연수 양식(수행평가 과제 및 채점기준표 개발) 모양의 워드 파일."""
    from urllib.parse import quote

    from server.item_docx import build_docx

    with closing(connect()) as conn:
        r = item_row(get_item_row(conn, iid))
    name = f"{r['data'].get('title') or '수행평가 과제'} (수행평가 과제 및 채점기준표).docx".replace("/", "·")
    return Response(
        build_docx({**r["data"], "subject": r["subject"]}),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"},
    )
