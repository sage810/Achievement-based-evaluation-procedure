"""문항 개발 → 워드 파일. 연수 양식 「[모듈 3] 수행평가 과제 및 채점기준표 개발(양식)」의 순서·칸을 따른다.

1️⃣ 수행평가 과제 정보표 → 2️⃣ 수행평가 문항 → 3️⃣ 채점기준표(성취수준 열 추가) → 4️⃣ 예시답안
"""
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LEVELS = ("A", "B", "C", "D", "E")
TEAL = "00B2B1"
HEAD_FILL = "D8ECED"
FONT = "맑은 고딕"


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(table) -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), TEAL)
        b.append(e)
    tblPr.append(b)


def write(cell, text: str, head: bool = False, center: bool = False) -> None:
    """칸에 글자를 넣는다. 줄바꿈(\\n)은 문단으로 나눈다."""
    cell.text = ""
    lines = (text or "").split("\n")
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.font.size = Pt(9.5)
        if head:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(TEAL)
    if head:
        shade(cell, HEAD_FILL)


def table(doc, rows: int, cols: int, widths: list[float]):
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    return t


def heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


def arrow_row(t, r: int) -> None:
    cell = t.cell(r, 0).merge(t.cell(r, len(t.columns) - 1))
    write(cell, "⇩", center=True)


def build_docx(d: dict) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(9.5)

    levels = d.get("levels") or {}
    traits = d.get("min_traits") or {}
    content = d.get("content") or {}
    std = d.get("standard", "")

    # ── 1️⃣ 수행평가 과제 정보표 ──
    heading(doc, "1️⃣ 수행평가 과제 정보표")
    t = table(doc, 6, 3, [4, 1, 12.4])
    write(t.cell(0, 0), "성취기준", head=True, center=True)
    write(t.cell(0, 1).merge(t.cell(0, 2)), "성취수준별 최소 능력자의 수행 특성", head=True, center=True)
    write(t.cell(1, 0).merge(t.cell(5, 0)), std)
    for i, lv in enumerate(LEVELS, 1):
        write(t.cell(i, 1), lv, head=True, center=True)
        write(t.cell(i, 2), traits.get(lv, ""))
    doc.add_paragraph()

    n = 18  # 과목·과제명·성취수준 머리 1+5·분석·⇩·범주 1+3·⇩·평가 요소·⇩·평가 방법
    t = table(doc, n, 3, [4, 1, 12.4])
    r = 0
    write(t.cell(r, 0), "과목 · 영역", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), f"{d.get('subject', '')} · {d.get('area', '')}".strip(" ·"))
    r += 1
    write(t.cell(r, 0), "평가 과제명", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), d.get("title", ""))
    r += 1
    write(t.cell(r, 0), "성취기준", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), "성취기준별 성취수준", head=True, center=True)
    r += 1
    write(t.cell(r, 0).merge(t.cell(r + 4, 0)), std)
    for i, lv in enumerate(LEVELS):
        write(t.cell(r + i, 1), lv, head=True, center=True)
        write(t.cell(r + i, 2), levels.get(lv, ""))
    r += 5
    write(t.cell(r, 0), "교육과정 분석", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), d.get("curriculum_analysis", ""))
    r += 1
    arrow_row(t, r)
    r += 1
    write(t.cell(r, 0), "범주", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), "내용 요소", head=True, center=True)
    r += 1
    for key, label in (("knowledge", "지식·이해"), ("skills", "과정·기능"), ("attitude", "가치·태도")):
        write(t.cell(r, 0), label, head=True, center=True)
        write(t.cell(r, 1).merge(t.cell(r, 2)), content.get(key, ""))
        r += 1
    arrow_row(t, r)
    r += 1
    write(t.cell(r, 0), "평가 요소", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), d.get("elements", ""))
    r += 1
    arrow_row(t, r)
    r += 1
    write(t.cell(r, 0), "평가 방법", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), d.get("method", ""))
    doc.add_paragraph()

    lessons = d.get("lessons") or []
    info = d.get("item_info") or []
    t = table(doc, 1 + len(lessons) + 2 + len(info), 3, [3, 8.2, 6.2])
    for c, h in enumerate(("단계", "교수･학습 활동", "평가 활동 및 주요 평가 내용")):
        write(t.cell(0, c), h, head=True, center=True)
    r = 1
    for ls in lessons:
        write(t.cell(r, 0), ls.get("stage", ""), head=True, center=True)
        write(t.cell(r, 1), ls.get("teaching", ""))
        write(t.cell(r, 2), ls.get("assessment", ""))
        r += 1
    write(t.cell(r, 0), "개발 방향 및\n활용 시 고려 사항", head=True, center=True)
    write(t.cell(r, 1).merge(t.cell(r, 2)), d.get("dev_direction", ""))
    r += 1
    for c, h in enumerate(("문항/과제 번호", "성취기준", "평가 요소")):
        write(t.cell(r, c), h, head=True, center=True)
    r += 1
    for it in info:
        write(t.cell(r, 0), it.get("no", ""), center=True)
        write(t.cell(r, 1), it.get("standard", ""))
        write(t.cell(r, 2), it.get("element", ""))
        r += 1

    # ── 2️⃣ 수행평가 문항 ──
    heading(doc, "2️⃣ 수행평가 문항")
    t = table(doc, 1, 1, [17.4])
    write(t.cell(0, 0), d.get("question", ""))

    # ── 3️⃣ 채점기준표 (성취수준 열 추가) ──
    heading(doc, "3️⃣ 채점기준표")
    rubric = d.get("rubric") or []
    total = sum(max(1, len(b.get("rows") or [])) for b in rubric)
    t = table(doc, 1 + total, 6, [1.4, 3, 3, 1.4, 7.4, 1.2])
    for c, h in enumerate(("문항", "평가 요소", "채점 요소", "성취수준", "수행 수준(채점기준)", "배점")):
        write(t.cell(0, c), h, head=True, center=True)
    r = 1
    for b in rubric:
        rows = b.get("rows") or [{}]
        last = r + len(rows) - 1
        write(t.cell(r, 0).merge(t.cell(last, 0)), b.get("no", ""), center=True)
        write(t.cell(r, 1).merge(t.cell(last, 1)), b.get("element", ""))
        write(t.cell(r, 2).merge(t.cell(last, 2)), b.get("criterion", ""))
        for i, row in enumerate(rows):
            write(t.cell(r + i, 3), row.get("level", ""), center=True)
            write(t.cell(r + i, 4), row.get("text", ""))
            write(t.cell(r + i, 5), str(row.get("score", "")), center=True)
        r = last + 1
    if d.get("basic_score"):
        # 「항목: 내용」 줄을 2열 표로 (같은 항목은 한 칸에)
        rows: list[list[str]] = []
        for line in filter(str.strip, d["basic_score"].split("\n")):
            label, sep, body = line.partition(":")
            if not sep or len(label.strip()) > 10:
                label, body = "기본점수", line
            label, body = label.strip(), body.strip()
            if rows and any(r[0] == label for r in rows):
                next(r for r in rows if r[0] == label)[1] += "\n" + body
            else:
                rows.append([label, body])
        doc.add_paragraph().paragraph_format.space_before = Pt(4)
        t = table(doc, len(rows), 2, [3.0, 14.4])
        for i, (label, body) in enumerate(rows):
            write(t.cell(i, 0), label, head=True, center=True)
            write(t.cell(i, 1), body)

    # ── 4️⃣ 예시답안 ──
    heading(doc, "4️⃣ 예시답안")
    answers = d.get("answers") or []
    t = table(doc, 1 + len(answers), 2, [2.2, 15.2])
    write(t.cell(0, 0), "문항번호", head=True, center=True)
    write(t.cell(0, 1), "예시 답안", head=True, center=True)
    for i, a in enumerate(answers, 1):
        write(t.cell(i, 0), a.get("no", ""), center=True)
        write(t.cell(i, 1), a.get("answer", ""))

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
